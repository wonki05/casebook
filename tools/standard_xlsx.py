"""Lossless XLSX standard-case catalog preparation and reviewed reconciliation.

This module never writes registry.json or toc.json. Use add_cases.py
standard-import to commit its validated plan. Workbook cells remain source
strings; reviewed citation corrections affect identity only.
"""
import copy
import hashlib
import json
import re
from collections import Counter, defaultdict
from difflib import SequenceMatcher
from pathlib import Path

from case_identity import identity, primary_case_number, decision_type, IDENTITY_FIELDS

SHEETS = {'민법': 'civil', '상법': 'commercial', '형법': 'criminal',
          '헌법': 'constitution', '행정법': 'administrative'}
EDITION = '2026'
CITATION = re.compile(r'^(.*?)\s+(?:(\d{4}-\d{2}-\d{2})\s+)?선고\s+(.+?)\s*$')


def normalized_title(text):
    return re.sub(r'[\s:：·ㆍ,()（）\[\]－—–-]+', '', text or '')


def read_workbook(path, reg, overrides=None):
    import openpyxl
    path = Path(path)
    overrides = overrides or {}
    row_overrides = overrides.get('rows', overrides)
    wb = openpyxl.load_workbook(path, data_only=True, read_only=False)
    missing = set(SHEETS) - set(wb.sheetnames)
    if missing:
        raise ValueError('필수 과목 시트 누락: ' + ', '.join(sorted(missing)))
    digest = hashlib.sha256(path.read_bytes()).hexdigest()
    try:
        source_file = path.resolve().relative_to(Path(__file__).resolve().parents[1]).as_posix()
    except ValueError:
        source_file = path.as_posix()
    source = {'file': source_file, 'sha256': digest, 'edition': EDITION}
    catalog = {'source': source, 'candidates': [], 'conflicts': [], 'subjects': {},
               'rawIdentities': [], 'sourceAnomalies': [],
               'excludedSheets': [s for s in wb.sheetnames if s not in SHEETS]}
    catalog['legacyIdentityCorrections'] = copy.deepcopy(overrides.get('legacyIdentityCorrections', {}))
    catalog['legacySourceExclusions'] = copy.deepcopy(overrides.get('legacySourceExclusions', {}))
    # Non-target procedure sheets still prove continued inclusion in this
    # edition. Reading their citations prevents a false deletion notice.
    for sheet in wb.sheetnames:
        if sheet in SHEETS or sheet == '개요':
            continue
        for row in range(2, wb[sheet].max_row + 1):
            match = CITATION.fullmatch(str(wb[sheet].cell(row, 3).value or ''))
            if match:
                court, date, number = match.groups()
                catalog['rawIdentities'].append({'subject': None, 'court': court, 'caseNumber': number,
                                                 'date': date, 'row': row, 'sheet': sheet})
    existing_by_key = defaultdict(list)
    for c in reg['cases']:
        existing_by_key[(c['court'], primary_case_number(c['caseNumber']), c['decisionDate'])].append(c)
    for sheet, subject in SHEETS.items():
        ws = wb[sheet]
        if [ws.cell(1, i).value for i in (1, 2, 3)] != ['번호', '제목', '판례 정보']:
            raise ValueError('예상하지 못한 시트 머리글: ' + sheet)
        count = 0
        for row in range(2, ws.max_row + 1):
            vals = [ws.cell(row, i).value for i in range(1, 7)]
            if not any(v is not None and str(v).strip() for v in vals):
                continue
            count += 1
            key = f'{sheet}:{row}'
            raw = {'sheet': sheet, 'row': row, 'subject': subject,
                   'title': vals[1], 'sourceCitation': vals[2], 'sourceOrder': vals[0]}
            patch = row_overrides.get(key) or {}
            if patch and not patch.get('evidence') and not patch.get('tocPath'):
                raise ValueError(f'{key}: 식별 보정의 evidence가 필요합니다')
            match = CITATION.fullmatch(str(vals[2] or ''))
            if match:
                court, date, number = match.groups()
                raw.update(court=court, date=date, caseNumber=number)
                catalog['rawIdentities'].append({'subject': subject, 'court': court, 'caseNumber': number, 'row': row, 'sheet': sheet})
            else:
                catalog['conflicts'].append({**raw, 'reason': '판례 인용 형식을 해석할 수 없음'})
                continue
            candidate = dict(raw)
            for field in ('court', 'date', 'caseNumber', 'decisionType', 'enBanc', 'tocPath', 'matchExistingId', 'title', 'lawgo'):
                if field in patch:
                    candidate[field] = copy.deepcopy(patch[field])
            if patch.get('evidence'):
                candidate['identityReview'] = {'sourceCitation': vals[2], 'evidence': copy.deepcopy(patch['evidence']),
                                               'reason': patch.get('reason', '출처 대조에 따른 식별 정보 보완')}
            try:
                primary = primary_case_number(candidate['caseNumber'])
                matches = existing_by_key.get((candidate['court'], primary, candidate['date']), [])
                types = {c['decisionType'] for c in matches}
                if 'decisionType' not in candidate and len(types) == 1:
                    candidate['decisionType'] = next(iter(types))
                candidate['decisionType'] = decision_type(candidate)
                candidate.update(identity({**candidate, 'decisionDate': candidate['date']}))
                if not isinstance(candidate['title'], str) or not candidate['title'].strip():
                    raise ValueError('제목 누락')
            except (ValueError, TypeError) as exc:
                catalog['conflicts'].append({**raw, 'reason': str(exc)})
                continue
            std_source = {'verified': True, 'book': '표준판례', 'edition': EDITION,
                          'file': source_file, 'format': 'xlsx', 'sheet': sheet,
                          'rows': [row], 'sha256': digest, 'preferred': True}
            fields = {3: 'facts' if ws.cell(1, 4).value == '사실관계의 요지' else 'issues', 5: 'holdings'}
            for col, name in fields.items():
                if name in patch.get('excludeFields', []):
                    if not patch.get('evidence'):
                        raise ValueError(f'{key}: 필드 제외의 evidence가 필요합니다')
                    continue
                if vals[col] is not None and str(vals[col]).strip():
                    if not isinstance(vals[col], str):
                        raise ValueError(f'{key}: {name} 셀이 문자열이 아님')
                    std_source[name] = [{'n': '', 'text': vals[col]}]
            if any(name in std_source for name in ('issues', 'facts', 'holdings')):
                candidate['standardSource'] = std_source
            if 'selectionReason' not in patch.get('excludeFields', []):
                candidate['selectionReason'] = vals[4] if vals[4] is not None else ''
            candidate['standard'] = {'edition': EDITION, 'entryId': key, 'order': count, 'removed': False,
                                     'file': source_file, 'sheet': sheet, 'row': row,
                                     'sourceOrder': vals[0], 'sourceCitation': vals[2]}
            if ws.cell(row, 3).hyperlink:
                candidate['standard']['sourceUrl'] = ws.cell(row, 3).hyperlink.target
            if patch.get('excludeFields') or ('title' in patch and patch['title'] != vals[1]):
                anomaly = {'sheet': sheet, 'row': row, 'subject': subject,
                           'caseNumber': candidate['caseNumber'], 'title': candidate['title'],
                           'sourceTitle': vals[1], 'sourceCitation': vals[2],
                           'excludedFields': patch.get('excludeFields', []),
                           'sourceValues': {name: vals[col] for col, name in (fields | {4: 'selectionReason'}).items() if name in patch.get('excludeFields', [])},
                           'reason': patch.get('reason'), 'evidence': patch.get('evidence')}
                candidate['standard']['sourceReview'] = copy.deepcopy(anomaly)
                catalog['sourceAnomalies'].append(anomaly)
            candidate['status'] = 'ok'
            catalog['candidates'].append(candidate)
        if count < 1:
            raise ValueError('빈 과목 시트: ' + sheet)
        catalog['subjects'][subject] = {'sheet': sheet, 'rows': count}
    wb.close()
    return catalog


def choose_existing(candidate, records, used):
    explicit = candidate.get('matchExistingId')
    if explicit:
        matches = [c for c in records if c['id'] == explicit]
        if len(matches) != 1 or explicit in used or not candidate.get('identityReview'):
            raise ValueError('검토한 기존 배치 연결이 유효하지 않음: ' + explicit)
        return matches[0]
    options = [c for c in records if c['subject'] == candidate['subject'] and
               c['decisionId'] == candidate['decisionId'] and c['id'] not in used]
    if not options:
        return None
    def rank(c):
        meta = c.get('standard') or {}
        return (meta.get('entryId') == candidate['standard']['entryId'],
                normalized_title(c['title']) == normalized_title(candidate['title']),
                c.get('tocPath') == candidate.get('tocPath'),
                SequenceMatcher(None, normalized_title(c['title']), normalized_title(candidate['title'])).ratio())
    return max(options, key=rank)


def assign_toc(catalog, reg, toc, reviewed_paths=None):
    """Use retained placements, reviewed row mapping, or agreeing adjacent anchors.

    A boundary with different adjacent chapter paths stays unresolved; do not
    turn a plausible topical guess into source metadata.
    """
    reviewed_paths = reviewed_paths or {}
    standards = [c for c in reg['cases'] if c['group'] == 'standard']
    used = set()
    unresolved = []
    for candidate in catalog['candidates']:
        old = choose_existing(candidate, standards, used)
        if old:
            used.add(old['id'])
            candidate['matchedId'] = old['id']
        key = candidate['standard']['entryId']
        mapping = reviewed_paths.get(key)
        if mapping:
            candidate['tocPath'] = list(mapping['tocPath'] if isinstance(mapping, dict) else mapping)
            candidate['tocMapping'] = 'reviewed-row'
        elif candidate.get('tocPath'):
            candidate['tocMapping'] = 'reviewed-identity'
        elif old:
            candidate['tocPath'] = list(old['tocPath'])
            candidate['tocMapping'] = 'existing-placement'
    for subject in catalog['subjects']:
        items = [c for c in catalog['candidates'] if c['subject'] == subject]
        for i, c in enumerate(items):
            if c.get('tocPath'):
                continue
            prev = next((x for x in reversed(items[:i]) if x.get('tocPath')), None)
            nxt = next((x for x in items[i+1:] if x.get('tocPath')), None)
            if prev and nxt and prev['tocPath'] == nxt['tocPath']:
                c['tocPath'] = list(prev['tocPath'])
                c['tocMapping'] = 'agreeing-adjacent-anchors'
            else:
                unresolved.append({k: c.get(k) for k in ('sheet', 'row', 'subject', 'title', 'caseNumber', 'date')} | {
                    'previous': {k: prev.get(k) for k in ('row', 'title', 'tocPath')} if prev else None,
                    'next': {k: nxt.get(k) for k in ('row', 'title', 'tocPath')} if nxt else None})
    catalog['unresolvedToc'] = unresolved
    return catalog


def plan_import(reg, toc, catalog, make_entry):
    """Return a fully validated immutable-input import plan; write only via CLI."""
    if catalog.get('conflicts') or catalog.get('unresolvedToc'):
        raise ValueError(f"미해결 식별 {len(catalog.get('conflicts', []))}건 / 목차 {len(catalog.get('unresolvedToc', []))}건")
    updated = copy.deepcopy(reg)
    standards = [c for c in updated['cases'] if c['group'] == 'standard']
    initial_decisions = {c['decisionId'] for c in reg['cases']}
    initial_standard = {c['decisionId'] for c in standards}
    # An explicitly corrected erroneous cell is not evidence that its wrong
    # docket remains selected. Keep raw citations for audit, use reviewed
    # effective identities for target sheets and raw procedure-sheet listings.
    comparison_rows = [{**c, 'sheet': c['standard']['sheet']} for c in catalog['candidates']]
    comparison_rows.extend(c for c in catalog['rawIdentities'] if c.get('subject') is None)
    raw_keys = {(c['court'], primary_case_number(c['caseNumber'])) for c in comparison_rows}
    current_decisions = {c['decisionId'] for c in catalog['candidates']}
    locations = defaultdict(set)
    for raw in comparison_rows:
        locations[(raw['court'], primary_case_number(raw['caseNumber']))].add(raw['sheet'])
    by_subject = defaultdict(set)
    for c in catalog['candidates']:
        by_subject[c['subject']].add(c['decisionId'])
    report = {'source': copy.deepcopy(catalog['source']), 'added': [], 'removed': [],
              'changed': [], 'matches': [], 'identityCorrections': [], 'conflicts': [],
              'unresolvedToc': [], 'retainedPlacements': [], 'summary': {}}
    report['sourceAnomalies'] = copy.deepcopy(catalog.get('sourceAnomalies', []))
    report['relocatedPlacements'] = []
    report['reviewedCorrections'] = [{k: copy.deepcopy(c.get(k)) for k in ('subject', 'title', 'sheet', 'row', 'caseNumber', 'date', 'decisionId', 'identityReview')}
                                    for c in catalog['candidates'] if c.get('identityReview')]
    used = set()
    def info(c):
        meta = c.get('standard') or {}
        return {k: copy.deepcopy(c.get(k)) for k in ('id', 'subject', 'title', 'caseNumber', 'court', 'decisionId', 'tocPath')} | {
            'date': c.get('decisionDate'), 'sheet': meta.get('sheet'), 'row': meta.get('row'),
            'entryId': meta.get('entryId'), 'order': meta.get('order')}
    for cid, correction in catalog.get('legacyIdentityCorrections', {}).items():
        old = next((c for c in standards if c['id'] == cid), None)
        if not old or not correction.get('evidence'):
            raise ValueError('기존 식별 보정에는 표준판례 ID와 검증 근거가 필요합니다: ' + cid)
        before = {k: old.get(k) for k in ('caseNumber', 'lookupKey', 'decisionDate') + IDENTITY_FIELDS}
        changed_record = dict(old)
        for field in ('caseNumber', 'decisionDate', 'decisionType', 'court'):
            if field in correction:
                changed_record[field] = correction[field]
        changed_record['lookupKey'] = primary_case_number(changed_record['caseNumber'])
        for field in ('caseNumberId', 'courtId'):
            changed_record.pop(field, None)
        changed_record.update(identity(changed_record))
        if old['decisionId'] != changed_record['decisionId']:
            aliases = updated.setdefault('identityAliases', {})
            aliases[old['decisionId']] = changed_record['decisionId']
            aliases[cid] = changed_record['decisionId']
            legacy = old.setdefault('legacyIds', [])
            if old['decisionId'] not in legacy:
                legacy.append(old['decisionId'])
            review = {'previousIdentity': before, 'evidence': correction['evidence'], 'reason': correction.get('reason')}
            old.setdefault('identityHistory', []).append(review)
            old.update({k: changed_record[k] for k in ('caseNumber', 'lookupKey', 'decisionDate', 'court') + IDENTITY_FIELDS})
            report['identityCorrections'].append(info(old) | {'review': review})
    for cid, exclusion in catalog.get('legacySourceExclusions', {}).items():
        old = next((c for c in standards if c['id'] == cid), None)
        if not old or not exclusion.get('evidence'):
            raise ValueError('기존 표준판례 출처 제외에는 ID와 검증 근거가 필요합니다: ' + cid)
        source = old.get('standardSource')
        if not source:
            continue
        removed_fields = [field for field in exclusion.get('fields', []) if field in source]
        if removed_fields:
            original = copy.deepcopy(source)
            history = old.setdefault('standardSourceHistory', [])
            if original not in history:
                history.append(original)
            for field in removed_fields:
                source.pop(field)
            old['standardSourceReview'] = {'excludedFields': removed_fields,
                'evidence': copy.deepcopy(exclusion['evidence']), 'reason': exclusion.get('reason')}
            report['sourceAnomalies'].append(info(old) | {'excludedFields': removed_fields,
                'sourceValues': {field: original[field] for field in removed_fields},
                'reason': exclusion.get('reason'), 'evidence': exclusion['evidence']})
    for candidate in catalog['candidates']:
        old = choose_existing(candidate, standards, used)
        entry = make_entry(updated, toc, 'standard', candidate['subject'], candidate['caseNumber'],
                           candidate['date'], candidate['title'], candidate['tocPath'],
                           court=candidate['court'], en_banc=candidate.get('enBanc', old.get('enBanc') if old else None),
                           decision_type=candidate['decisionType'], extra={
                               'standardSource': candidate.get('standardSource'), 'standard': candidate['standard'],
                               'selectionReason': candidate.get('selectionReason'), 'lawgo': candidate.get('lawgo')})
        if old:
            used.add(old['id'])
            before = copy.deepcopy(old)
            if old['decisionId'] != entry['decisionId']:
                if not candidate.get('identityReview') or candidate.get('matchExistingId') != old['id']:
                    raise ValueError('검토 없이 기존 재판 식별을 변경할 수 없음: ' + old['id'])
                review = copy.deepcopy(candidate['identityReview'])
                review['previousIdentity'] = {k: old.get(k) for k in ('caseNumber', 'lookupKey', 'decisionDate') + IDENTITY_FIELDS}
                old.setdefault('identityHistory', []).append(review)
                aliases = updated.setdefault('identityAliases', {})
                aliases[old['decisionId']] = entry['decisionId']
                aliases[old['id']] = entry['decisionId']
                legacy = old.setdefault('legacyIds', [])
                if old['decisionId'] not in legacy:
                    legacy.append(old['decisionId'])
                old.update({k: entry[k] for k in ('caseNumber', 'lookupKey', 'decisionDate', 'court') + IDENTITY_FIELDS})
                report['identityCorrections'].append({**info(old), 'review': review})
            previous_source = old.get('standardSource')
            if previous_source and entry.get('standardSource') and previous_source != entry['standardSource']:
                history = old.setdefault('standardSourceHistory', [])
                if previous_source not in history:
                    history.append(copy.deepcopy(previous_source))
            if not old.get('standard'):
                old['standardPrevious'] = {'title': old['title'], 'tocPath': list(old['tocPath']),
                                          'registryOrder': next(i for i, x in enumerate(reg['cases'], 1) if x['id'] == old['id'])}
            for field in ('title', 'tocPath', 'standardSource', 'standard', 'selectionReason', 'enBanc', 'lawgo'):
                if field in entry:
                    old[field] = copy.deepcopy(entry[field])
            if candidate.get('identityReview'):
                old['standard']['identityReview'] = copy.deepcopy(candidate['identityReview'])
            record = info(old)
            record['previousTitle'] = before['title']
            record['previousTocPath'] = before['tocPath']
            report['matches'].append(record)
            if old != before:
                report['changed'].append(record)
        else:
            updated['cases'].append(entry)
            used.add(entry['id'])
            classification = 'additional-placement' if entry['decisionId'] in initial_standard else ('new-standard' if entry['decisionId'] in initial_decisions else 'new-decision')
            report['added'].append(info(entry) | {'classification': classification})
    subject_order = defaultdict(int)
    for c in standards:
        subject_order[c['subject']] += 1
        if c['id'] in used:
            continue
        present = c['decisionId'] in current_decisions
        # A same-docket but different-date row is an identity discrepancy, not
        # evidence that this legacy judgment was removed from the edition.
        raw_key = (c['court'], primary_case_number(c['caseNumber']))
        docket_present = raw_key in raw_keys
        removed = not present and not docket_present
        c.setdefault('standardPrevious', {'title': c['title'], 'tocPath': list(c['tocPath']), 'registryOrder': next(i for i, x in enumerate(reg['cases'], 1) if x['id'] == c['id'])})
        c['standard'] = {'edition': EDITION, 'entryId': 'legacy:' + c['id'],
                         'order': catalog['subjects'][c['subject']]['rows'] + subject_order[c['subject']],
                         'removed': removed, 'file': catalog['source']['file'],
                         'comparison': 'absent-from-workbook' if removed else 'retained-legacy-placement'}
        if removed:
            report['removed'].append(info(c) | {'reason': '전체 과목 시트에 동일 재판 및 사건번호가 없음'})
        else:
            sheets = sorted(locations.get(raw_key, set()))
            own_sheet = catalog['subjects'][c['subject']]['sheet']
            reason = '동일 재판 또는 사건번호가 계속 수록됨; 기존 배치를 삭제하지 않음'
            if sheets and own_sheet not in sheets:
                reason = '다른 과목 시트에 수록됨: ' + ', '.join(sheets)
                report['relocatedPlacements'].append(info(c) | {'reason': reason, 'currentSheets': sheets})
            report['retainedPlacements'].append(info(c) | {'reason': reason, 'currentSheets': sheets})
    for subject, meta in catalog['subjects'].items():
        report['summary'][subject] = {**meta, 'addedPlacements': sum(c['subject'] == subject for c in report['added']),
            'removedPlacements': sum(c['subject'] == subject for c in report['removed']),
            'matchedPlacements': sum(c['subject'] == subject for c in report['matches']),
            'retainedPlacements': sum(c['subject'] == subject for c in report['retainedPlacements'])}
    return updated, report
